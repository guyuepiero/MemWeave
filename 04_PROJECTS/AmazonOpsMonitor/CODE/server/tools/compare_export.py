"""同类目跨时间对比 → 导出 Excel 明细（独立进程执行，使用 SCRAPER_PY env 的 openpyxl）

用法:
    python compare_export.py <json_a> <json_b> <out_dir> [out_name.xlsx]

产出 sheet:
    1 大盘概览   2 榜单换血   3 排名变动   4 品牌变化   5 明细全量
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import compare as C  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
SUB_FONT = Font(bold=True, size=10, color="1F4E79")
NOTE_FONT = Font(size=9, color="808080")


def _sheet(wb, title, headers, rows, widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = NOTE_FONT
        r0 = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            cell = ws.cell(i, j, v)
            cell.font = Font(size=9.5)
            cell.alignment = Alignment(vertical="top", wrap_text=len(str(v or "")) > 60)
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1).coordinate if rows else None
    return ws


def _md(v):
    return "—" if v is None else v


def build(out: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    m = out["market"]
    da, db = out["a"]["date"], out["b"]["date"]

    def d1(d):
        return None if not d or d.get("delta_pct") is None else d

    def arrow(d):
        if not d or d.get("delta_pct") is None:
            return "数据不足"
        p = d["delta_pct"]
        return ("↑" if p > 0 else "↓" if p < 0 else "→") + f" {abs(p):.1f}%"

    # ---- 1 大盘概览 ----
    sd, rd = d1(m["sales_same"]), d1(m["revenue_same"])
    sa, ra = d1(m["sales_all"]), d1(m["revenue_all"])
    ap = d1(m["avg_price_same"])
    rv = m["review_same"]
    rows = [
        ["对比区间", f"{da}  →  {db}", f"间隔 {m['days']} 天", "", ""],
        ["同批在榜 ASIN", m["same_n"], f"占两期 {round(m['same_n']/max(m['n_a'],m['n_b'])*100)}%", "", ""],
        ["大盘月销量(估·同批口径)", _md(sd and round(sd['old'])), _md(sd and round(sd['new'])),
         _md(sd and round(sd['delta'])), arrow(sd)],
        ["大盘月销量(估·全榜含换血)", _md(sa and round(sa['old'])), _md(sa and round(sa['new'])),
         _md(sa and round(sa['delta'])), arrow(sa)],
        ["大盘月销售额(估·同批)", _md(rd and round(rd['old'])), _md(rd and round(rd['new'])),
         _md(rd and round(rd['delta'])), arrow(rd)],
        ["大盘月销售额(估·全榜含换血)", _md(ra and round(ra['old'])), _md(ra and round(ra['new'])),
         _md(ra and round(ra['delta'])), arrow(ra)],
        ["同批商品均价", _md(ap and f"${ap['old']:.2f}"), _md(ap and f"${ap['new']:.2f}"),
         _md(ap and f"{'+' if (ap['delta'] or 0) >= 0 else ''}{ap['delta']:.2f}"), arrow(ap)],
        ["同批评论存量", rv["old"], rv["new"], f"+{rv['inc']}", f"{rv['per_day']} 条/天"],
        ["Top10 重合", m["top10_overlap"], "/10", "", ""],
        ["Top30 重合", m["top30_overlap"], "/30", "", ""],
        ["掉出/新进", m["drop_n"], m["new_n"], "", ""],
        ["上升/下降/持平", out["rankings"]["up"], out["rankings"]["down"], out["rankings"]["flat"],
         f"平均位移 {_md(out['rankings']['avg_abs_move'])} 位"],
    ]
    note = ("口径说明：Amazon 无官方绝对销量，月销量/月销售额为页面档位文本(如 1K+/500+)按档中值=下限×1.5 的估算值，"
            "仅供同口径相对趋势参考。大盘主口径=两期均在榜同批 ASIN，剔除榜单换血噪声。")
    _sheet(wb, "1大盘概览", ["指标", da, db, "增减", "趋势"], rows,
           widths=[26, 16, 16, 14, 14], note=note)

    # 价格带 / 月销档分布（竖排便于阅读）
    pb = []
    for i, x in enumerate(m["price_bands_a"]):
        y = m["price_bands_b"][i]
        pb.append(["价格带", x["label"], x["n"], y["n"], y["n"] - x["n"]])
    sb = []
    for i, x in enumerate(m["sales_bands_a"]):
        y = m["sales_bands_b"][i]
        sb.append(["月销档(区间)", x["label"], x["n"], y["n"], y["n"] - x["n"]])
    _sheet(wb, "1b价格与销量带", ["类型", "档位", da, db, "增减"], pb + [["", "", "", "", ""]] + sb,
           widths=[12, 12, 8, 8, 8])

    # ---- 2 榜单换血 ----
    def churn_rows(items, status, date_col):
        out_rows = []
        for x in items:
            out_rows.append([status, x["asin"], x["brand"], x.get("rank"), date_col,
                             _md(x["price_raw"]), _md(x["sales_raw"]), _md(x["review"]), _md(x["rating"]),
                             "✓" if x.get("ac") else "", "✓" if x.get("bb") else "", x["title"]])
        return out_rows

    rows2 = churn_rows(out["churn"]["dropped"], "掉出", da) + churn_rows(out["churn"]["entered"], "新进", db)
    _sheet(wb, "2榜单换血", ["状态", "ASIN", "品牌", "名次", "所在期", "价格", "月销", "评论数", "评分", "AC标", "BS标", "标题"],
           rows2, widths=[6, 12, 12, 6, 11, 9, 8, 9, 6, 6, 6, 60], note=f"{da}→{db}：掉出 {len(out['churn']['dropped'])} · 新进 {len(out['churn']['entered'])}")

    # ---- 3 排名变动 ----
    rows3 = []
    for x in out["rankings"]["all"]:
        d = x["delta"]
        rows3.append([x["asin"], x["brand"], x["rank_old"], x["rank_new"],
                      ("↑" + str(d)) if d > 0 else (("↓" + str(-d)) if d < 0 else "—"),
                      _md(x["price_old"] and round(x["price_old"], 2)), _md(x["price_new"] and round(x["price_new"], 2)),
                      _md(x["sales_old"]), _md(x["sales_new"]), _md(x["rev_inc"]),
                      _md(x["rating_old"]), _md(x["rating_new"]), x["title"]])
    _sheet(wb, "3排名变动", ["ASIN", "品牌", "旧名次", "新名次", "变动", "旧价$", "新价$", "旧月销", "新月销", "评论增量", "旧评分", "新评分", "标题"],
           rows3, widths=[12, 12, 7, 7, 7, 8, 8, 8, 8, 9, 6, 6, 55], note=f"同批 {len(rows3)} 个：升 {out['rankings']['up']} · 降 {out['rankings']['down']} · 平 {out['rankings']['flat']}")

    # ---- 4 品牌变化 ----
    rows4 = []
    for x in out["brands"]["moves"]:
        rows4.append([x["brand"], x["cnt_a"], x["cnt_b"], x["cnt_delta"],
                      f"{x['share_a']}%", f"{x['share_b']}%", _md(x["top_a"]), _md(x["top_b"]),
                      _md(x["avg_rank_a"] and round(x["avg_rank_a"], 1)), _md(x["avg_rank_b"] and round(x["avg_rank_b"], 1))])
    for x in out["brands"]["new"]:
        rows4.append([f"★新进 {x['brand']}", 0, x["cnt"], x["cnt"], "0%", f"{round(x['cnt']/max(out['b']['n'],1)*100,1)}%", "", x["top"], "", ""])
    for x in out["brands"]["gone"]:
        rows4.append([f"✕掉榜 {x['brand']}", x["cnt"], 0, -x["cnt"], f"{round(x['cnt']/max(out['a']['n'],1)*100,1)}%", "0%", x["top"], "", "", ""])
    _sheet(wb, "4品牌变化", ["品牌", f"上榜数{da}", f"上榜数{db}", "变动", f"份额{da}", f"份额{db}", f"最佳名次{da}", f"最佳名次{db}", f"均名次{da}", f"均名次{db}"],
           rows4, widths=[18, 10, 10, 8, 9, 9, 10, 10, 10, 10], note="份额=该品牌 ASIN 数/榜单100；均名次=该品牌在榜商品平均名次(越小越靠前)")

    # ---- 5 明细全量 ----
    rows5 = []
    for x in out["detail"]:
        st = {"kept": "在榜", "dropped": "掉出", "entered": "新进"}[x["status"]]
        rows5.append([st, x["asin"], x["brand"], _md(x.get("rank_old")), _md(x.get("rank_new")),
                      _md(x.get("delta")), _md(x.get("price_old") and round(x["price_old"], 2)),
                      _md(x.get("price_new") and round(x["price_new"], 2)),
                      _md(x.get("rev_old")), _md(x.get("rev_new")), _md(x.get("rev_inc")),
                      _md(x.get("sales_old")), _md(x.get("sales_new")), _md(x.get("rating_old")), _md(x.get("rating_new")),
                      x["title"]])
    _sheet(wb, "5明细全量", ["状态", "ASIN", "品牌", "旧名次", "新名次", "名次变动", "旧价$", "新价$", "旧评论", "新评论", "评论增量", "旧月销", "新月销", "旧评分", "新评分", "标题"],
           rows5, widths=[6, 12, 12, 7, 7, 8, 8, 8, 9, 9, 9, 8, 8, 6, 6, 55], note=f"{da}→{db} 全量 {len(rows5)} 条：在榜 {out['churn']['kept']} · 掉出 {len(out['churn']['dropped'])} · 新进 {len(out['churn']['entered'])}")
    return wb


def main():
    if len(sys.argv) < 4:
        print("用法: python compare_export.py <json_a> <json_b> <out_dir> [out_name.xlsx]")
        return 2
    pa, pb, out_dir = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    da = json.loads(pa.read_text(encoding="utf-8"))
    db = json.loads(pb.read_text(encoding="utf-8"))
    out = C.compare(da, db)
    if "error" in out:
        print("ERROR:", out["error"])
        return 1
    name = sys.argv[4] if len(sys.argv) > 4 else f"compare_{out['a']['date']}_vs_{out['b']['date']}.xlsx"
    dest = out_dir / name
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = build(out)
    wb.save(dest)
    print("OK", dest)
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())
